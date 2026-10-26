from pathlib import Path
import re,zipfile,hashlib,json
from lxml import etree
from docx import Document
root=Path(r"D:\桌面\single cell")
out=root/"修订交付_2026-10-26_v2"
final=out/"manuscripts_1026_revised_final_2026-10-26_final.docx"
original=root/"manuscripts_1026改.docx"

def zipstats(p):
    with zipfile.ZipFile(p) as z:
        assert z.testzip() is None
        xml=etree.fromstring(z.read('word/document.xml'))
        media=[n for n in z.namelist() if n.startswith('word/media/')]
        rels=etree.fromstring(z.read('word/_rels/document.xml.rels'))
        return len(media),len(xml.xpath("//*[local-name()='drawing']")),len(rels)
o=zipstats(original); f=zipstats(final)
doc=Document(final); paras=doc.paragraphs
ref_start=next(i for i,p in enumerate(paras) if p.text.strip().startswith('References:'))
refs=[p for p in paras if p.style.name=='EndNote Bibliography']
heads=[p.text for p in paras if re.match(r'^3\.[1-9] ',p.text.strip())]
figs=[p.text for p in paras if re.match(r'^Figure [1-9]\.',p.text.strip())]
text='\n'.join(p.text for p in paras)
first=[]; bad=[]
for p in paras[:ref_start]:
    if 'Acknowledgements' in p.text: break
    for m in re.finditer(r'\(([0-9,\s]+)\)',p.text):
        nums=[int(x) for x in re.findall(r'\d+',m.group(1))]
        for n in nums:
            if 1<=n<=len(refs) and n not in first: first.append(n)
            elif n>len(refs): bad.append(n)
checks={
 'zip_and_relationships':True,
 'media_unchanged':o[:2]==f[:2],
 'result_sections_6':len(heads)==6 and not any(h.startswith('3.7 ') for h in heads),
 'figure_captions_6':len(figs)==6,
 'references_41':len(refs)==41,
 'citation_first_order_1_to_41':first==list(range(1,42)),
 'no_out_of_range_citations':not bad,
 'no_placeholders':'{{cite:' not in text,
 'no_double_parentheses':all(f'(({n}' not in text for n in range(10)),
 'no_S19':'Table S19' not in text,
 'table_S7_S18_cited':all(f'Table S{n}' in text for n in range(7,19)),
 'fig5F_trajectory':'Fig. 5F' in text and 'Candidate state-differentiation trajectory' in text,
 'fig5G_composition':'Fig. 5G' in text and '(G) Cell-state composition' in text,
 'github_in_methods':'Code availability. All analysis scripts' in text and 'https://github.com/jiangwy2026/hepatopancreatic-sinus-snrna-analysis' in text,
 'discussion_expanded':len([p for p in paras if p.text.startswith(('Circulating phagocytes','The Hs niche','At single-nucleus','The dissociation','The external LyO','The ECM-integrin','The evolutionary','The syncytium','The main boundary','Taken together'))])==10,
 'reference_numbering_style':all(p._p.pPr is not None and p._p.pPr.numPr is not None for p in refs),
}
critical=['zip_and_relationships','media_unchanged','result_sections_6','figure_captions_6','references_41','citation_first_order_1_to_41','no_out_of_range_citations','no_placeholders','no_double_parentheses','no_S19','table_S7_S18_cited','fig5F_trajectory','fig5G_composition','github_in_methods','discussion_expanded','reference_numbering_style']
assert all(checks[k] for k in critical),checks
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
report=['# Final manuscript validation','',f'- Reference count: {len(refs)}',f'- First-citation order: {first[:5]} ... {first[-5:]}',f'- Result sections: {len(heads)}',f'- Main figure captions: {len(figs)}',f'- Embedded media: original {o[0]}, final {f[0]}',f'- Drawings: original {o[1]}, final {f[1]}',f'- SHA256: `{sha(final)}`','', '## Checks','']
report += [f'- {k}: {"PASS" if v else "FAIL"}' for k,v in checks.items()]
(out/'最终文档校验_2026-10-26.md').write_text('\n'.join(report),encoding='utf-8')
print(json.dumps(checks,indent=2)); print('saved',out/'最终文档校验_2026-10-26.md')