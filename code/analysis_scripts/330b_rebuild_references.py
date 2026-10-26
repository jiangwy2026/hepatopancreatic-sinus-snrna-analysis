# -*- coding: utf-8 -*-
"""Convert citations to stable keys, rebuild reference order, and insert final references."""
from pathlib import Path
from copy import deepcopy
import re, json, csv
from docx import Document
from docx.text.paragraph import Paragraph

ROOT=Path(r"D:\桌面\single cell")
V2=ROOT/"修订交付_2026-10-26_v2"
SRC=V2/"work"/"manuscript_stage_content_2026-10-26.docx"
DST=V2/"manuscripts_1026_revised_final_2026-10-26_final.docx"
doc=Document(SRC)

old_refs=[p.text.strip() for p in doc.paragraphs if p.style.name=='EndNote Bibliography']
assert len(old_refs)==40, len(old_refs)
ref_text={f"ref{i:02d}":t.replace("Elife.","eLife.") for i,t in enumerate(old_refs,1)}
# Correct the official issue metadata for Sun et al.
ref_text['ref30']="Sun M, Li S, Liu Y, Li F. Single-nucleus RNA sequencing illuminates a functional analog of lymphoid organ in crustacean. J Adv Res. 2026 Aug;86:305-320. doi: 10.1016/j.jare.2025.11.018."
# Remove two references whose topics did not support the rewritten Discussion.
for k in ['ref31','ref32']: ref_text.pop(k,None)

new_refs={
'ghersi2023':"Ghersi JJ, Baldissera G, Hintzen J, Luff SA, Cheng S, Xia IF, Sturgeon CM, Nicoli S. Haematopoietic stem and progenitor cell heterogeneity is inherited from the embryonic endothelium. Nat Cell Biol. 2023 Aug;25(8):1135-1145. doi: 10.1038/s41556-023-01187-9.",
'gao2022':"Gao S, Shi Q, Zhang Y, Liang G, Kang Z, Huang B, Ma D, Wang L, Jiao J, Fang X, Xu CR, Liu L, Xu X, Göttgens B, Li C, Liu F. Identification of HSC/MPP expansion units in fetal liver by single-cell spatiotemporal transcriptomics. Cell Res. 2022 Jan;32(1):38-53. doi: 10.1038/s41422-021-00540-7.",
'donnelly2024':"Donnelly H, Ross E, Xiao Y, Hermantara R, Taqi AF, Doherty-Boyd WS, Cassels J, Tsimbouri PM, Dunn KM, Hay J, Cheng A, Meek RMD, Jain N, West C, Wheadon H, Michie AM, Peault B, West AG, Salmeron-Sanchez M, Dalby MJ. Bioengineered niches that recreate physiological extracellular matrix organisation to support long-term haematopoietic stem cells. Nat Commun. 2024 Jul 10;15(1):5791. doi: 10.1038/s41467-024-50054-0.",
}
ref_text.update(new_refs)

# Convert existing numeric citations to stable keys before renumbering.
ref_heading_idx=next(i for i,p in enumerate(doc.paragraphs) if p.text.strip().startswith('References:'))
body_paras=doc.paragraphs[:ref_heading_idx]
sup_map={'⁰':0,'¹':1,'²':2,'³':3,'⁴':4,'⁵':5,'⁶':6,'⁷':7,'⁸':8,'⁹':9}

def replace_run_text(p,new):
    rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for r in list(p.runs): r._r.getparent().remove(r._r)
    run=p.add_run(new)
    if rpr is not None and run._r.rPr is None: run._r.insert(0,deepcopy(rpr))

def expand_group(s):
    out=[]
    for part in re.split(r'\s*,\s*',s):
        if re.fullmatch(r'\d{1,2}\s*[-–—]\s*\d{1,2}',part):
            a,b=map(int,re.split(r'\s*[-–—]\s*',part)); out.extend(range(a,b+1))
        elif re.fullmatch(r'\d{1,2}',part): out.append(int(part))
        else: return None
    return out

for p in body_paras:
    t=p.text
    # Parenthetical numeric citations.
    def paren_sub(m):
        nums=expand_group(m.group(1))
        if not nums or any(n<1 or n>40 for n in nums): return m.group(0)
        return '{{cite:'+','.join(f'ref{n:02d}' for n in nums)+'}}'
    nt=re.sub(r'\((\d{1,2}(?:\s*[,–—-]\s*\d{1,2})*)\)',paren_sub,t)
    # Unicode superscript citations in the Introduction.
    def sup_sub(m):
        n=int(''.join(str(sup_map[c]) for c in m.group(0)))
        return f'{{{{cite:ref{n:02d}}}}}' if 1<=n<=40 else m.group(0)
    nt=re.sub(r'[⁰¹²³⁴⁵⁶⁷⁸⁹]+',sup_sub,nt)
    if nt!=t: replace_run_text(p,nt)

# Assign final numbers by first appearance.
body_paras=doc.paragraphs[:next(i for i,p in enumerate(doc.paragraphs) if p.text.strip().startswith('References:'))]
order=[]; seen=set()
for p in body_paras:
    for m in re.finditer(r'\{\{cite:([^}]+)\}\}',p.text):
        for key in [x.strip() for x in m.group(1).split(',')]:
            if key not in ref_text: raise KeyError(f'Unknown reference key {key}')
            if key not in seen: seen.add(key); order.append(key)
if not order: raise RuntimeError('No citations found')
missing=[k for k in ref_text if k not in seen]
if missing: raise RuntimeError('Uncited references remaining: '+','.join(missing))
num={k:i for i,k in enumerate(order,1)}
for p in body_paras:
    t=p.text
    def cite_sub(m):
        nums=[num[k.strip()] for k in m.group(1).split(',')]
        return '('+', '.join(map(str,sorted(set(nums))))+')'
    nt=re.sub(r'\{\{cite:([^}]+)\}\}',cite_sub,t)
    if nt!=t: replace_run_text(p,nt)

# Build final reference paragraphs in citation order.
ref_template=next(p for p in doc.paragraphs if p.style.name=='EndNote Bibliography')
template_p=deepcopy(ref_template._p)
template_rpr=None
if ref_template.runs and ref_template.runs[0]._r.rPr is not None: template_rpr=deepcopy(ref_template.runs[0]._r.rPr)
for p in [p for p in doc.paragraphs if p.style.name=='EndNote Bibliography']:
    p._p.getparent().remove(p._p)
anchor=next(p for p in doc.paragraphs if p.text.strip().startswith('References:'))
for key in order:
    new_p=deepcopy(template_p); anchor._p.addnext(new_p); q=Paragraph(new_p,anchor._parent)
    for r in list(q.runs): r._r.getparent().remove(r._r)
    run=q.add_run(ref_text[key])
    if template_rpr is not None: run._r.insert(0,deepcopy(template_rpr))
    anchor=q

# Save mapping and final document.
rows=[]
for key in order:
    n=num[key]; doi=(re.search(r'10\.\d{4,9}/[^\s]+',ref_text[key]).group(0).rstrip('.;,') if re.search(r'10\.\d{4,9}/[^\s]+',ref_text[key]) else '')
    rows.append({'number':n,'key':key,'doi':doi,'reference':ref_text[key]})
with (V2/'work'/'ref_audit'/'final_reference_map.csv').open('w',newline='',encoding='utf-8-sig') as f:
    w=csv.DictWriter(f,fieldnames=['number','key','doi','reference']); w.writeheader(); w.writerows(rows)
(V2/'work'/'ref_audit'/'final_reference_map.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
doc.save(DST)
print('saved',DST,'references',len(order))
print('first 12',[(r['number'],r['key']) for r in rows[:12]])
print('last 12',[(r['number'],r['key']) for r in rows[-12:]])
