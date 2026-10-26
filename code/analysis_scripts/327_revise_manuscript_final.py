# -*- coding: utf-8 -*-
"""Revise the Hs manuscript: six result sections, integrated former 3.7, revised Fig. 5/6 legends."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph

ROOT=Path(r"D:\桌面\single cell")
SRC=ROOT/"manuscripts_1026改.docx"
OUT=ROOT/"修订交付_2026-10-26_v2"
DST=OUT/"manuscripts_1026_revised_final_2026-10-26_noS19.docx"
doc=Document(SRC)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise KeyError(prefix)

def settext(p,text):
    rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run=p.add_run(text)
    if rpr is not None and run._r.rPr is None:
        run._r.insert(0,deepcopy(rpr))

def append(p,suffix):
    if suffix.strip() not in p.text:
        settext(p,p.text.rstrip()+suffix)

# Locate the seven result blocks before moving XML elements.
p92=find("3.1 "); p95=find("3.2 "); p100=find("3.3 "); p103=find("3.4 ")
p107=find("3.5 "); p108=find("To dissect the cellular composition"); p110=find("3.6 "); p113=find("3.7 ")
p109=find("After hemospasia, Hs cellular composition"); p111=find("Cell cycle analysis showed")
p112=find("After hemospasia, the increase in total PH proportion")
p114=find("To test whether the progenitor-like identity of PH")
p115=find("Next, it was necessary to exclude")
p116=find("Since the progenitor identity of Hs PH")
p117=find("Comparison before and after hemospasia")

# Reorder former 3.7 material into the two preceding sections:
# 3.5 receives the candidate niche communication analysis; 3.6 receives external identity and cell-cycle adjustment.
p109._p.addnext(p117._p)
p109._p.addnext(p116._p)
p112._p.addnext(p115._p)
p112._p.addnext(p114._p)
p113._p.getparent().remove(p113._p)

# Six result headings, aligned verbatim with the six figure captions.
settext(p92,"3.1 The hepatopancreatic sinus is the major cell proliferation center after hemospasia")
settext(p95,"3.2 Hs hemolymph has distinct metabolomic, proteomic, and cellular characteristics from systemic hemolymph")
settext(p100,"3.3 Functional classification and proliferative activity validation of differentially expressed metabolites in Hs plasma")
settext(p103,"3.4 Classification and dynamic analysis of differentially expressed proteins in Hs plasma after hemospasia")
settext(p107,"3.5 Single-cell atlas of Hs cell states, state-specific marker programs, and PH-to-GH transcriptional ordering after hemospasia")
settext(p110,"3.6 Hemospasia remodels a proliferative, progenitor-like PH pool with continuous transcriptional progression")

# Results text after 3.5: use nuclei consistently, correct panel letters and new table numbers.
settext(p108,"To dissect the cellular composition of the hepatopancreatic sinus (Hs) and its response to hemospasia, we performed single-nucleus RNA sequencing on Hs tissue from control and 6-h hemospasia groups. A total of 19,795 nuclear barcodes were identified, and 9,658 nuclei were retained after doublet and quality filtering (5,349 control and 4,309 hemospasia; Table S8, Table S9). t-SNE separated the nuclei into 10 transcriptional states: PH, GH, TS, MH2, F, R, B, M, MyoEpi, and PM-like (Fig. 5C,D), each with a distinguishable marker-gene profile (Table S10). On the basis of marker genes, reference-atlas matching, and label-transfer results, these states grouped into blood cell/immune-related lineages and hepatopancreas-like cells.")
settext(p109,"After hemospasia, the Hs nuclear composition underwent state-selective remodeling (Fig. 5F; Table S12). PH increased from 6.8% to 15.7% (2.30-fold, FDR = 9.0 x 10^-44), R increased from 2.2% to 5.1% (2.29-fold, FDR = 6.1 x 10^-14), and MH2 increased from 5.8% to 12.6% (2.17-fold, FDR = 3.9 x 10^-31). In contrast, PM-like decreased from 10.9% to 2.9% (0.27-fold, FDR = 4.1 x 10^-49), and TS decreased from 52.1% to 40.6% (0.78-fold, FDR = 5.4 x 10^-29). These changes identify PH and associated immune states as the principal mobilizable populations at 6 h after hemospasia.")

# Former 3.7 communication paragraphs now follow the composition result in section 3.5.
settext(p116,"We then examined whether this state reorganization coincided with changes in candidate niche communication (Fig. S2A-D; Table S16, Table S17). In the ECM-integrin axis, MyoEpi->PH and MyoEpi->GH had the highest communication strengths, 0.24 and 0.25, respectively, compared with 0.14 for TS, 0.15 for M, and 0.11 for MH2. In the Notch axis, M->PH and M->GH had communication strengths of 0.02 and 0.03, respectively. At the individual ligand-receptor level, VWF->ITGB1 contributed most to MyoEpi->PH/GH communication (0.178 and 0.181; state-label permutation P < 0.001). THBS4->ITGB1 was 0.068 and 0.070 for M->PH and M->GH (P = 0.034 and 0.026), JAG1->NOTCH1 was 0.015 and 0.026 (P = 0.007 and 0.001), and POSTN->ITGB1 reached 0.006 for M->GH (P = 0.009). ITGB1 expression was sparse in PH and GH nuclei (positive fractions 7.6% and 8.4%; median expression 0; Table S18). The ECM-integrin result therefore defines a candidate communication axis rather than confirmed high-abundance receptor activation.")
settext(p117,"Comparison of the two conditions showed a directional rearrangement of candidate sources. The VWF->ITGB1 probability from MyoEpi to PH increased from 0.129 to 0.287 (log2FC = 1.16, 500-permutation P = 0.028), while the corresponding MyoEpi->GH change was 0.153 to 0.263 (log2FC = 0.78, P = 0.068). M-to-PH and M-to-GH estimates decreased, but the pair-level changes were not significant under the same 500-permutation test (P = 0.60 and 0.35 for THBS4->ITGB1, respectively). Thus, the candidate signal shifted from M toward MyoEpi sources in parallel with PH expansion, while the nucleus-level evidence remains exploratory because each condition contains one pooled library (Table S16).")

# 3.6: cell-cycle-independent PH state and continuous substate progression.
settext(p111,"Cell-cycle analysis showed that PH nuclei had high proliferative potential: their G2/M fraction was 46.2% (n = 1,044), the highest among well-populated states, whereas the G2/M fractions of GH, TS, MH2, F, R, B, and PM-like ranged from 25.0% to 33.9% (Fig. 6A). M nuclei had a G2/M fraction of 53.3%, but comprised only 45 nuclei. PH nuclei were further resolved into four transcriptional substates, s0-s3 (Fig. 6B; Table S11): s0 was ECM/ribosome-high, s1 was progenitor-like, s2 was immune/ECM-primed, and s3 was TGM1-low. The substates had distinguishable marker combinations: s0 was enriched for PEN-3 and ECM-related transcripts; s1 expressed progenitor-like markers including Igfbp7 and ccnd2; s2 was enriched for THBS4, TGM1, and LOC113810024; and s3 expressed LOC113818748, LOC113823090, and LOC113823089 with low TGM1 expression (Fig. S1A; Table S11).")
settext(p112,"After hemospasia, total PH expansion was accompanied by a shift in substate composition: s0 decreased from 48.1% to 12.4%, s1 increased from 24.9% to 38.3%, s2 increased from 7.9% to 31.7%, and s3 changed from 19.1% to 17.6%. Genes related to progenitor maintenance, ECM organization, and cell maturation changed continuously along PH pseudotime rather than forming fully separated expression blocks (Fig. 6C; Table S14). In the PH-only t-SNE, s1 localized mainly on the left, s3 mainly on the right, s0 in the center, and s2 in an intermediate transition region (Fig. 6D). The s1->s0->s2->s3 line connects substate centroids to summarize transcriptional continuity and does not represent lineage tracing. Across states, PH occupied the earlier end of the projected pseudotime distribution (median 0.278; IQR 0.228-0.328) and GH the later end (median 0.386; IQR 0.368-0.413; Fig. 5G). This places the intra-PH substate progression within a PH-to-GH transcriptional ordering while leaving lineage direction unresolved. State-level control-versus-hemospasia DEGs are reported in Table S13.")

# Former 3.7 identity paragraphs now conclude section 3.6.
settext(p114,"To test whether the progenitor-like identity of PH extended beyond Hs, we scored the external LyO progenitor top-50 program across states (Fig. 6E; Fig. S1C,D; Table S15). The score was highest in PH (mean UCell-style score 0.525), with 77% of PH nuclei above the global median; the remaining states had mean scores of 0.469-0.501, with 20%-60% of nuclei above the global median. In the overlap analysis, the LyO progenitor program and the PH top-50 marker set shared four genes (66.9-fold enrichment, P = 3.7 x 10^-7, BH-adjusted P = 1.9 x 10^-6), whereas the other examined state marker sets had zero overlap. This is program-level concordance between PH and the external LyO progenitor program.")
settext(p115,"To separate progenitor identity from cell-cycle activity, we compared the PH2 identity score with the cell-cycle score across Hs nuclei (Fig. S1B). PH nuclei were generally above the background cell-cycle-identity trend, and the mean adjusted identity residual was 1.14 z units (95% CI 1.10-1.18; Fig. 6F). We then compared Hs PH with hemolymph PH. The raw PH2 score was higher in hemolymph PH (0.887 versus 0.840), but its cell-cycle percentile was also much higher (0.975 versus 0.568). After cell-cycle adjustment, the residual was 1.14 (95% CI 1.10-1.18) for Hs PH and -0.13 (95% CI -0.15 to -0.10) for hemolymph PH, a difference of 1.27 (95% CI 1.22-1.31; Fig. 6F). Hs PH therefore retains a stronger progenitor-identity program beyond the cell-cycle background than hemolymph PH.")

# Methods: connect the manuscript to the new S7-S18 tables without renumbering S1-S6.
append(find("2.10 Hemospasia, nuclei isolation")," Library and sample metadata are provided in Table S7.")
append(find("Cell Ranger 8.0.0 was run")," Nucleus-level metadata and QC summaries are provided in Table S8, Table S9.")
append(find("Marker genes for the 10 states")," Significant marker genes are listed in Table S10.")
append(find("State composition was summarized descriptively")," Composition source data and exploratory tests are reported in Table S12.")
append(find("PH nuclei were subset and reclustered")," All significant PH substate markers are listed in Table S11.")
append(find("A PH–GH trajectory was constructed")," Pseudotime-bin source data are provided in Table S14.")
append(find("External progenitor programs comprised")," Program-level scores and overlap tests are provided in Table S15.")
append(find("To separate cell-cycle activity from progenitor identity")," Cross-organ summary values are displayed in Fig. 6F.")
append(find("Candidate niche communication was analyzed")," Candidate communication and orthology mappings are provided in Table S16, Table S17.")
append(find("Hemospasia-induced changes in communication")," ITGB1 summaries are provided in Table S18.")
# Table S19 is deferred by user request.

# Figure captions: exact same core titles as the six result headings.
settext(find("Figure 1."),"Figure 1. The hepatopancreatic sinus is the major cell proliferation center after hemospasia.")
settext(find("Figure 2."),"Figure 2. Hs hemolymph has distinct metabolomic, proteomic, and cellular characteristics from systemic hemolymph.")
settext(find("Figure 3."),"Figure 3. Functional classification and proliferative activity validation of differentially expressed metabolites in Hs plasma.")
settext(find("Figure 4."),"Figure 4. Classification and dynamic analysis of differentially expressed proteins in Hs plasma after hemospasia.")
settext(find("Figure. 5."),"Figure 5. Single-cell atlas of Hs cell states, state-specific marker programs, and PH-to-GH transcriptional ordering after hemospasia.")
settext(find("Figure. 6."),"Figure 6. Hemospasia remodels a proliferative, progenitor-like PH pool with continuous transcriptional progression.")

# Figure 5 legend, including the previously missing panel G.
settext(find("(A) Experimental workflow for single-cell RNA sequencing of hepatopancreatic sinus"),"(A) Experimental workflow for single-nucleus RNA sequencing of Hs tissue from control and 6-h hemospasia groups.")
settext(find("(B) Quality-control metrics for retained cells"),"(B) Quality-control metrics for 9,658 retained nuclei (5,349 control and 4,309 hemospasia). Nucleus-level metadata and QC summaries are provided in Table S8, Table S9.")
settext(find("(C) t-SNE visualization of all Hs cells"),"(C) t-SNE visualization of all 9,658 Hs nuclei, identifying 10 transcriptional states: PH, GH, TS, MH2, F, R, B, M, MyoEpi, and PM-like. State marker evidence is provided in Table S10.")
settext(find("(D) Heat map of marker-gene expression"),"(D) Heat map of marker-gene expression across the 10 states, with the right inset showing markers distinguishing the MyoEpi and M populations.")
settext(find("(E) t-SNE distributions of control and hemospasia cells"),"(E) Control and hemospasia nuclei plotted on identical t-SNE coordinates (5,349 and 4,309 nuclei, respectively).")
settext(find("(F) Cell-state composition and log"),"(F) Cell-state composition and log2 fold change in the hemospasia group relative to control. Differences were assessed using two-proportion z tests with Benjamini-Hochberg correction. Hemospasia induced state-selective reorganization, including expansion of PH, R, and MH2 and contraction of PM-like and TS populations. Counts and exploratory tests are provided in Table S12.")
# Insert G after F.
p_f=find("(F) Cell-state composition and log")
new_p=deepcopy(p_f._p); p_f._p.addnext(new_p)
gp=Paragraph(new_p,p_f._parent)
for r in list(gp.runs): r._r.getparent().remove(r._r)
run=gp.add_run("(G) PH-to-GH transcriptional ordering. Left, PH and GH nuclei on the Hs t-SNE colored by diffusion pseudotime. Middle, pseudotime distributions for PH and GH. Right, row-standardized expression of representative PH marker genes across 20 equal-count pseudotime bins. DPT describes relative transcriptional order and does not establish lineage; source data are provided in Table S14.")
if p_f.runs and p_f.runs[0]._r.rPr is not None: run._r.insert(0,deepcopy(p_f.runs[0]._r.rPr))

# Figure 6 legend.
settext(find("(A) Cell-cycle composition of the 10 Hs states"),"(A) Cell-cycle composition of the 10 Hs states, showing G1, S, and G2M fractions. PH had the highest G2M fraction among well-populated states (46.2%; n = 1,044). M had a G2M fraction of 53.3% but contained only 45 nuclei.")
settext(find("(B) Composition of PH substates"),"(B) Composition of PH substates s0-s3 in control and hemospasia groups. s0, ECM/ribosome-high; s1, progenitor-like; s2, immune/ECM-primed; s3, TGM1-low. Hemospasia increased s1 and s2 while reducing s0. Marker evidence is provided in Table S11.")
settext(find("(C) Dynamics of representative PH marker genes"),"(C) Dynamics of representative PH marker genes along pseudotime, displayed as row-standardized expression from early to late pseudotime. Pseudotime-bin source data are provided in Table S14.")
settext(find("(D) PH-only t-SNE plots"),"(D) PH-only t-SNE plots showing all PH nuclei and control and hemospasia PH nuclei separately. The s1->s0->s2->s3 line connects substate centroids and summarizes transcriptional continuity; it does not represent lineage tracing.")
settext(find("(E) UCell scores for the external LyO progenitor"),"(E) UCell scores for the external LyO progenitor and hemolymph prohemocyte programs projected onto the Hs t-SNE, with state-level summaries. Both programs reached their highest mean scores in PH. Program-level scores and overlap tests are provided in Table S15.")
settext(find("(F) Relationship between cell-cycle and progenitor-identity"),"(F) Relationship between cell-cycle and progenitor-identity programs in Hs PH and hemolymph PH, together with adjusted progenitor-identity residuals after subtraction of the cell-cycle background trend. Hs PH retained a higher progenitor-identity program after cell-cycle adjustment.")

# Minor consistency fixes in earlier text.
for p in doc.paragraphs:
    t=p.text
    if "Fig. 1E， Table S1" in t:
        settext(p,t.replace("Fig. 1E， Table S1","Fig. 1E, Table S1"))
    if "Fig. 1E, Table S1, Fig. S1-S10" in t:
        settext(p,t.replace("Fig. 1E, Table S1, Fig. S1-S10","Fig. 1E and Table S1"))

doc.save(DST)
print("saved",DST)
