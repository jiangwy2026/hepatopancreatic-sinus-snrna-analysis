# -*- coding: utf-8 -*-
import json, math
from pathlib import Path
from collections import Counter
import openpyxl
import pandas as pd

root=Path(r"D:\桌面\single cell\修订交付_2026-10-26_v2")
payload=json.loads((root/"source_tables"/"supplementary_tables.json").read_text(encoding="utf-8"))
t={x["sheet"]:x for x in payload["tables"]}
issues=[]

# Cross-table identity and count checks.
s8=t["S8"]["rows"]; s9=t["S9"]["rows"]; s12=t["S12"]["rows"]
assert len(s8)==9658
assert len({r[0] for r in s8})==9658
c8=Counter(r[1] for r in s8)
assert c8=={"Hs_control":5349,"Hs_hemospasia":4309},c8
for r in s9: assert c8[r[0]]==r[2],(r[0],r[2],c8[r[0]])
assert sum(r[1] for r in s12)==5349 and sum(r[3] for r in s12)==4309
assert Counter(r[9] for r in s8)=={"TS":4533,"GH":1462,"PH":1044,"MH2":857,"PM-like":710,"F":377,"R":338,"B":227,"MyoEpi":65,"M":45}
assert Counter(r[10] for r in s8 if r[10])=={"PH_s0":260,"PH_s1":351,"PH_s2":244,"PH_s3":189}
assert Counter(r[11] for r in s8)=={"G1":1,"S":1,"G2M":1} or set(r[11] for r in s8)<= {"G1","S","G2M"}
assert sum(r[12] is not None and math.isfinite(r[12]) for r in s8)==2506
assert sum(r[13] is not None and math.isfinite(r[13]) for r in s8)==9658
assert all(r[14] is not None and r[15] is not None and r[16] is not None and r[17] is not None for r in s8)
assert sum(r[9]=="PH" for r in s8)==1044
assert all(0 <= r[7] <= 1 for r in s9)
assert sum(int(r[9].split("n_nuclei=")[1].split(";")[0]) for r in t["S14"]["rows"])==1044
assert len([r for r in t["S14"]["rows"] if r[2]=="control"])==20
assert len([r for r in t["S14"]["rows"] if r[2]=="hemospasia"])==20
for sheet in ["S10","S11","S13"]:
    for r in t[sheet]["rows"]:
        pidx=5 if sheet in {"S10","S11"} else 5
        assert 0 <= r[pidx] <= 1 and 0 <= r[pidx-1] <= 1
for r in t["S16"]["rows"]:
    assert r[9] is None or 0 <= r[9] <= 1
# All S7-S19 labels are present in the manuscript.
docx_text='\n'.join(__import__('docx').Document(root/"manuscripts_1026_revised_final_2026-10-26.docx").paragraphs[i].text for i in range(len(__import__('docx').Document(root/"manuscripts_1026_revised_final_2026-10-26.docx").paragraphs)))
for n in range(7,20):
    if f"Table S{n}" not in docx_text: issues.append(f"manuscript lacks Table S{n}")

# Workbook row/header integrity and source-value spot checks.
wb=openpyxl.load_workbook(root/"Supplementary_Tables_S7-S19_2026-10-26.xlsx",read_only=True,data_only=True)
for sheet,spec in t.items():
    ws=wb[sheet]
    it=ws.iter_rows(values_only=True)
    rows=list(it)
    header=rows[4][:spec["n_cols"]]
    data=[list(r[:spec["n_cols"]]) for r in rows[5:5+spec["n_rows"]]]
    if list(header)!=spec["headers"]: issues.append(f"{sheet}: header mismatch")
    if len(data)!=spec["n_rows"]: issues.append(f"{sheet}: workbook rows {len(data)} vs {spec['n_rows']}")
    for idx in [0, min(1,len(data)-1), len(data)-1]:
        if idx>=0:
            src=spec["rows"][idx]; got=data[idx]
            for a,b in zip(src,got):
                if isinstance(a,(int,float)) and isinstance(b,(int,float)):
                    if not math.isclose(float(a),float(b),rel_tol=1e-6,abs_tol=1e-9): issues.append(f"{sheet}: numeric mismatch row {idx}"); break
                elif a!=b:
                    # openpyxl may normalize None/empty strings
                    if not ((a in (None,"") and b in (None,""))): issues.append(f"{sheet}: value mismatch row {idx}: {a!r} vs {b!r}"); break
wb.close()
print("ISSUES",issues)
print("table rows",[(x["sheet"],x["n_rows"],x["n_cols"]) for x in payload["tables"]])
(root/"work"/"final_validation.txt").write_text(("PASS\n" if not issues else "\n".join(issues)+"\n"),encoding="utf-8")
