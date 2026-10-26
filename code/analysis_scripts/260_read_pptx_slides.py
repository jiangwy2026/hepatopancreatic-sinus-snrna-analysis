from pptx import Presentation
import sys
sys.stdout.reconfigure(encoding="utf-8")
p=r"D:\桌面\single cell\图片-1020.pptx"
prs=Presentation(p)
print("slides",len(prs.slides))
for i,slide in enumerate(prs.slides,1):
    if i not in (6,7,8,9): continue
    print("\n"+"="*90)
    print("SLIDE",i)
    for sh in slide.shapes:
        t=""
        if sh.has_text_frame:
            t=" | ".join(x.text.strip() for x in sh.text_frame.paragraphs if x.text.strip())
        if t: print(f"[TEXT] {t[:1200]}")
        if sh.has_table:
            print("[TABLE]")
            for r in sh.table.rows:
                print("   ", " | ".join(c.text.strip() for c in r.cells)[:500])
        if sh.shape_type==13:
            print(f"[PICTURE] name={sh.name} size={sh.width/914400:.2f}x{sh.height/914400:.2f}in")
    if slide.has_notes_slide:
        n=slide.notes_slide.notes_text_frame.text.strip()
        if n: print("[NOTES]",n[:1500])
