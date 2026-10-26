# -*- coding: utf-8 -*-
"""Apply revised Results, Discussion, figure legends and Method statements."""
from pathlib import Path
import sys
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
sys.path.insert(0,str(Path(__file__).parent))
from importlib import import_module
R=import_module('330_content_results')
D=import_module('330_content_discussion')

ROOT=Path(r"D:\桌面\single cell")
SRC=ROOT/"manuscripts_1026改.docx"
OUT=ROOT/"修订交付_2026-10-26_v2"
STAGE=OUT/"work"/"manuscript_stage_content_2026-10-26.docx"
doc=Document(SRC)

def find(prefix):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix): return p
    raise KeyError(prefix)

def settext(p,text):
    rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for r in list(p.runs): r._r.getparent().remove(r._r)
    run=p.add_run(text)
    if rpr is not None and run._r.rPr is None: run._r.insert(0,deepcopy(rpr))
    return p

def append(p,suffix):
    if suffix.strip() not in p.text: settext(p,p.text.rstrip()+suffix)

def insert_after(p,text):
    new_p=deepcopy(p._p); p._p.addnext(new_p); q=Paragraph(new_p,p._parent)
    for r in list(q.runs): r._r.getparent().remove(r._r)
    run=q.add_run(text)
    if p.runs and p.runs[0]._r.rPr is not None: run._r.insert(0,deepcopy(p.runs[0]._r.rPr))
    return q

def insert_discussion_after(anchor,text,model):
    new_p=deepcopy(model._p); anchor._p.addnext(new_p); q=Paragraph(new_p,model._parent)
    for r in list(q.runs): r._r.getparent().remove(r._r)
    run=q.add_run(text)
    if model.runs and model.runs[0]._r.rPr is not None: run._r.insert(0,deepcopy(model.runs[0]._r.rPr))
    return q

# Locate result blocks.
p92=find("3.1 "); p95=find("3.2 "); p100=find("3.3 "); p103=find("3.4 ")
p107=find("3.5 "); p108=find("To dissect the cellular composition"); p110=find("3.6 "); p113=find("3.7 ")
p109=find("After hemospasia, Hs cellular composition"); p111=find("Cell cycle analysis showed")
p112=find("After hemospasia, the increase in total PH proportion")
p114=find("To test whether the progenitor-like identity of PH")
p115=find("Next, it was necessary to exclude")
p116=find("Since the progenitor identity of Hs PH")
p117=find("Comparison before and after hemospasia")

# Integrate former 3.7 material.
p109._p.addnext(p117._p); p109._p.addnext(p116._p)
p112._p.addnext(p115._p); p112._p.addnext(p114._p)
p113._p.getparent().remove(p113._p)

for prefix,text in R.RESULT_HEADINGS.items(): settext(find(prefix),text)
settext(p108,R.RESULT_PARAGRAPHS['p108'])
settext(p109,R.RESULT_PARAGRAPHS['p109'])
traj=insert_after(p109,R.RESULT_PARAGRAPHS['trajectory'])
settext(p116,R.RESULT_PARAGRAPHS['comm1'])
settext(p117,R.RESULT_PARAGRAPHS['comm2'])
insert_after(p117,R.RESULT_PARAGRAPHS['bridge'])
settext(p111,R.RESULT_PARAGRAPHS['p111'])
settext(p112,R.RESULT_PARAGRAPHS['p112'])
settext(p114,R.RESULT_PARAGRAPHS['external'])
settext(p115,R.RESULT_PARAGRAPHS['identity'])

# Methods: table pointers and public code availability.
append(find("2.10 Hemospasia, nuclei isolation")," Library and sample metadata are provided in Table S7.")
append(find("Cell Ranger 8.0.0 was run")," Nucleus-level metadata and QC summaries are provided in Table S8, Table S9.")
append(find("Marker genes for the 10 states")," Significant marker genes are listed in Table S10.")
append(find("State composition was summarized descriptively")," Composition source data and exploratory tests are reported in Table S12.")
append(find("PH nuclei were subset and reclustered")," All significant PH substate markers are listed in Table S11.")
append(find("A PH–GH trajectory was constructed")," Pseudotime source data are provided in Table S14.")
append(find("External progenitor programs comprised")," Program-level scores and overlap tests are provided in Table S15.")
append(find("To separate cell-cycle activity from progenitor identity")," Cross-organ summary values are displayed in Fig. 6F.")
append(find("Candidate niche communication was analyzed")," Candidate communication and orthology mappings are provided in Table S16, Table S17.")
append(find("Hemospasia-induced changes in communication")," ITGB1 summaries are provided in Table S18.")
pstat=find("Statistical tests were two-sided")
pcode=insert_after(pstat,"Code availability. All analysis scripts and supplementary source data are publicly available at https://github.com/jiangwy2026/hepatopancreatic-sinus-snrna-analysis (release v2026.10.26-final).")

# Figure captions and legends.
settext(find("Figure 1."),"Figure 1. The hepatopancreatic sinus is the major cell proliferation center after hemospasia.")
settext(find("Figure 2."),"Figure 2. Hs hemolymph has distinct metabolomic, proteomic, and cellular characteristics from systemic hemolymph.")
settext(find("Figure 3."),"Figure 3. Functional classification and proliferative activity validation of differentially expressed metabolites in Hs plasma.")
settext(find("Figure 4."),"Figure 4. Classification and dynamic analysis of differentially expressed proteins in Hs plasma after hemospasia.")
settext(find("Figure. 5."),"Figure 5. Single-nucleus transcriptomic landscape of Hs blood-cell states and a candidate differentiation trajectory after hemospasia.")
settext(find("Figure. 6."),"Figure 6. Hemospasia selectively expands a progenitor-like PH pool with substate reprogramming and identity beyond the cell cycle.")
settext(find("(A) Experimental workflow for single-cell RNA sequencing"),"(A) Experimental workflow for single-nucleus RNA sequencing of Hs tissue from control and 6-h hemospasia groups.")
settext(find("(B) Quality-control metrics for retained cells"),"(B) Quality-control metrics for 9,658 retained nuclei (5,349 control and 4,309 hemospasia). Nucleus-level metadata and QC summaries are provided in Table S8, Table S9.")
settext(find("(C) t-SNE visualization of all Hs cells"),"(C) t-SNE visualization of all 9,658 Hs nuclei, identifying 10 transcriptional states: PH, GH, TS, MH2, F, R, B, M, MyoEpi, and PM-like. State marker evidence is provided in Table S10.")
settext(find("(D) Heat map of marker-gene expression"),"(D) Heat map of marker-gene expression across the 10 states, with the right inset showing markers distinguishing the MyoEpi and M populations.")
settext(find("(E) t-SNE distributions of control and hemospasia cells"),"(E) Control and hemospasia nuclei plotted on identical t-SNE coordinates (5,349 and 4,309 nuclei, respectively).")
settext(find("(F) Cell-state composition and log"),"(F) Candidate state-differentiation trajectory overlaid on the Hs t-SNE. Arrows connect PH, biosynthetic transition, GH and MH2 populations to summarize relative transcriptional transitions; they do not represent lineage tracing.")
pF=find("(F) Candidate state-differentiation trajectory")
insert_after(pF,"(G) Cell-state composition and log2 fold change in the hemospasia group relative to control. Differences were assessed using two-proportion z tests with Benjamini-Hochberg correction. Hemospasia induced state-selective reorganization, including expansion of PH, R and MH2 and contraction of PM-like and TS populations. Counts and exploratory tests are provided in Table S12.")

settext(find("(A) Cell-cycle composition of the 10 Hs states"),"(A) Cell-cycle composition of the 10 Hs states, showing G1, S and G2M fractions. PH had the highest G2M fraction among well-populated states (46.2%; n = 1,044). M had a G2M fraction of 53.3% but contained only 45 nuclei.")
settext(find("(B) Composition of PH substates"),"(B) Composition of PH substates s0-s3 in control and hemospasia groups. s0, ECM/ribosome-high; s1, progenitor-like; s2, immune/ECM-primed; s3, TGM1-low. Hemospasia increased s1 and s2 while reducing s0. Marker evidence is provided in Table S11.")
settext(find("(C) Dynamics of representative PH marker genes"),"(C) Dynamics of representative PH marker genes along pseudotime, displayed as row-standardized expression from early to late pseudotime. Pseudotime-bin source data are provided in Table S14.")
settext(find("(D) PH-only t-SNE plots"),"(D) PH-only t-SNE plots showing all PH nuclei and control and hemospasia PH nuclei separately. The s1-s0-s2-s3 line connects substate centroids and summarizes transcriptional continuity; it does not represent lineage tracing.")
settext(find("(E) UCell scores for the external LyO progenitor"),"(E) UCell scores for the external LyO progenitor and hemolymph prohemocyte programs projected onto the Hs t-SNE, with state-level summaries. Both programs reached their highest mean scores in PH. Program-level scores and overlap tests are provided in Table S15.")
settext(find("(F) Relationship between cell-cycle and progenitor-identity"),"(F) Relationship between cell-cycle and progenitor-identity programs in Hs PH and hemolymph PH, together with adjusted progenitor-identity residuals after subtraction of the cell-cycle background trend. Hs PH retained a higher progenitor-identity program after cell-cycle adjustment.")

# Replace Discussion.
disc=find("Discussion"); credit=find("CRediT authorship contribution statement")
paras=doc.paragraphs; start=next(i for i,p in enumerate(paras) if p.text.strip().startswith("Discussion")); end=next(i for i,p in enumerate(paras) if p.text.strip().startswith("CRediT authorship contribution statement"))
old_between=paras[start+1:end]
model=next((p for p in old_between if p.text.strip()),old_between[0] if old_between else disc)
for p in old_between: p._p.getparent().remove(p._p)
anchor=disc
for text in D.DISCUSSION_PARAGRAPHS:
    anchor=insert_discussion_after(anchor,text,model)

# Minor consistency fixes.
for p in doc.paragraphs:
    t=p.text
    if "Fig. 1E， Table S1" in t:
        settext(p,t.replace("Fig. 1E， Table S1","Fig. 1E, Table S1"))
    if "Fig. 1E, Table S1, Fig. S1-S10" in t:
        settext(p,t.replace("Fig. 1E, Table S1, Fig. S1-S10","Fig. 1E and Table S1"))

doc.save(STAGE)
print("stage saved",STAGE)
